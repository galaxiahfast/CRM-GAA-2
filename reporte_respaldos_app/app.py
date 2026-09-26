from __future__ import annotations
import json, os, re
from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path
from tkinter import END, BooleanVar, StringVar, Tk, filedialog, messagebox, simpledialog, ttk

APP_NAME="Generador de reportes de respaldos"

@dataclass
class ServerBackup:
    server:str; backup_time:str; missing:str=""; notes:str=""
    @property
    def names(self): return [x.strip() for x in re.split(r"[\n;]+",self.missing) if x.strip()]
    @property
    def status(self): return "Con incidencias" if self.names else "Realizado"

def summary(rows):
    affected=sum(bool(r.names) for r in rows)
    return {"servers":len(rows),"completed":len(rows)-affected,"affected":affected,"missing":sum(len(r.names) for r in rows)}

def safe_name(text): return re.sub(r"[^\w.-]+","_",text).strip("_.") or "reporte_respaldos"

def generate_docx(path,title,date,responsible,rows,notes):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches,Pt,RGBColor
    d=Document(); s=d.sections[0]; s.orientation=WD_ORIENT.LANDSCAPE; s.page_width,s.page_height=Inches(11),Inches(8.5); s.left_margin=s.right_margin=Inches(.7); s.top_margin=s.bottom_margin=Inches(.6)
    for n in ("Normal","Title"):
        st=d.styles[n]; st.font.name="Aptos"; st._element.rPr.rFonts.set(qn("w:ascii"),"Aptos"); st._element.rPr.rFonts.set(qn("w:hAnsi"),"Aptos"); st.font.color.rgb=RGBColor(0,0,0)
    d.styles["Normal"].font.size=Pt(9); d.styles["Title"].font.size=Pt(23); d.styles["Title"].font.bold=True
    d.add_paragraph(title,style="Title"); p=d.add_paragraph(); p.add_run(f"Fecha del reporte: {date}").bold=True
    if responsible:p.add_run(f"    Responsable: {responsible}")
    p.paragraph_format.space_after=Pt(12); sm=summary(rows); overall="SIN INCIDENCIAS" if not sm["affected"] else "REQUIERE ATENCIÓN"
    cards=d.add_table(rows=1,cols=5)
    for i,(cell,(lab,val)) in enumerate(zip(cards.rows[0].cells,(("Resultado",overall),("Servidores",sm["servers"]),("Completados",sm["completed"]),("Con incidencias",sm["affected"]),("No realizados",sm["missing"])))):
        sh=OxmlElement("w:shd"); sh.set(qn("w:fill"),("E7F4EC" if not sm["affected"] else "FCE8E6") if i==0 else "EAF2F8"); cell._tc.get_or_add_tcPr().append(sh); p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(lab.upper()+"\n").font.size=Pt(7); r=p.add_run(str(val)); r.bold=True; r.font.size=Pt(11)
    d.add_paragraph(); t=d.add_table(rows=1,cols=5); t.style="Table Grid"; heads=("Servidor","Último respaldo","Estado","Respaldos no realizados","Observaciones"); widths=(2.2,1.7,1.25,3.25,1.25)
    for c,h,w in zip(t.rows[0].cells,heads,widths):
        c.width=Inches(w); sh=OxmlElement("w:shd"); sh.set(qn("w:fill"),"17365D"); c._tc.get_or_add_tcPr().append(sh); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(h); r.bold=True;r.font.size=Pt(8);r.font.color.rgb=RGBColor(255,255,255)
    hdr=OxmlElement("w:tblHeader"); hdr.set(qn("w:val"),"true"); t.rows[0]._tr.get_or_add_trPr().append(hdr)
    for i,row in enumerate(rows,1):
        vals=(row.server,row.backup_time,row.status,"\n".join("• "+x for x in row.names) if row.names else "Ninguno",row.notes)
        for j,(c,v,w) in enumerate(zip(t.add_row().cells,vals,widths)):
            c.width=Inches(w); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i%2==0: sh=OxmlElement("w:shd"); sh.set(qn("w:fill"),"F4F7FA"); c._tc.get_or_add_tcPr().append(sh)
            p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if j in (1,2) else WD_ALIGN_PARAGRAPH.LEFT; r=p.add_run(v); r.font.size=Pt(8.5)
            if j==2:r.bold=True;r.font.color.rgb=RGBColor(185,43,39) if row.names else RGBColor(31,122,74)
    if notes: d.add_paragraph("Observaciones generales").runs[0].bold=True; d.add_paragraph(notes)
    d.save(path)

def generate_pdf(path,title,date,responsible,rows,notes):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import landscape,letter
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer
    st=getSampleStyleSheet(); cell=ParagraphStyle("cell",parent=st["BodyText"],fontSize=8,leading=10); center=ParagraphStyle("center",parent=cell,alignment=TA_CENTER); head=ParagraphStyle("head",parent=center,textColor=colors.white,fontName="Helvetica-Bold")
    def foot(c,d): c.saveState();c.setFont("Helvetica",7);c.setFillColor(colors.grey);c.drawRightString(10.3*inch,.3*inch,f"Reporte generado automáticamente  |  Página {d.page}");c.restoreState()
    doc=SimpleDocTemplate(str(path),pagesize=landscape(letter),leftMargin=.6*inch,rightMargin=.6*inch,topMargin=.45*inch,bottomMargin=.45*inch)
    story=[Paragraph(title,ParagraphStyle("title",parent=st["Title"],fontSize=22,textColor=colors.black,spaceAfter=7)),Paragraph(f"<b>Fecha del reporte:</b> {date}"+(f"&nbsp;&nbsp;&nbsp;&nbsp;<b>Responsable:</b> {responsible}" if responsible else ""),cell),Spacer(1,.12*inch)]
    sm=summary(rows); overall="SIN INCIDENCIAS" if not sm["affected"] else "REQUIERE ATENCIÓN"; cards=[[Paragraph(f"<b>{a}</b><br/><font size='11'><b>{b}</b></font>",center) for a,b in (("RESULTADO",overall),("SERVIDORES",sm["servers"]),("COMPLETADOS",sm["completed"]),("CON INCIDENCIAS",sm["affected"]),("NO REALIZADOS",sm["missing"]))]]
    ct=Table(cards,colWidths=[2.1*inch,1.35*inch,1.4*inch,1.5*inch,1.4*inch]);ct.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#EAF2F8")),("BACKGROUND",(0,0),(0,0),colors.HexColor("#E7F4EC") if not sm["affected"] else colors.HexColor("#FCE8E6")),("GRID",(0,0),(-1,-1),.5,colors.HexColor("#D9D9D9")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]));story += [ct,Spacer(1,.16*inch)]
    data=[[Paragraph(x,head) for x in ("Servidor","Último respaldo","Estado","Respaldos no realizados","Observaciones")]]
    for r in rows:data.append([Paragraph(r.server,cell),Paragraph(r.backup_time,center),Paragraph(f"<b>{r.status}</b>",center),Paragraph("<br/>".join("• "+x for x in r.names) if r.names else "Ninguno",cell),Paragraph(r.notes,cell)])
    tb=Table(data,colWidths=[2.15*inch,1.65*inch,1.25*inch,3.25*inch,1.2*inch],repeatRows=1);cmd=[("BACKGROUND",(0,0),(-1,0),colors.HexColor("#17365D")),("GRID",(0,0),(-1,-1),.45,colors.HexColor("#D9D9D9")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("PADDING",(0,0),(-1,-1),6)];[cmd.append(("BACKGROUND",(0,i),(-1,i),colors.HexColor("#F4F7FA"))) for i in range(2,len(data),2)];tb.setStyle(TableStyle(cmd));story.append(tb)
    if notes:story += [Spacer(1,.15*inch),Paragraph("<b>Observaciones generales</b>",cell),Paragraph(notes,cell)]
    doc.build(story,onFirstPage=foot,onLaterPages=foot)

class App:
    def __init__(self,root):
        self.root=root;root.title(APP_NAME);root.geometry("1080x680"); self.title=StringVar(value="Reporte de respaldos de servidores");self.date=StringVar(value=datetime.now().strftime("%d/%m/%Y"));self.resp=StringVar();self.server=StringVar();self.time=StringVar(value=datetime.now().strftime("%d/%m/%Y %I:%M %p"));self.note=StringVar();self.word=BooleanVar(value=True);self.pdf=BooleanVar(value=True);self.ui()
    def ui(self):
        o=ttk.Frame(self.root,padding=18);o.pack(fill="both",expand=True);ttk.Label(o,text=APP_NAME,font=("Segoe UI",18,"bold")).pack(anchor="w");ttk.Label(o,text="Una fila por servidor. Solo registra los nombres que no se respaldaron.").pack(anchor="w",pady=(2,12))
        m=ttk.LabelFrame(o,text="Datos del reporte",padding=10);m.pack(fill="x")
        for i,(lab,var) in enumerate((("Título",self.title),("Fecha",self.date),("Responsable",self.resp))):m.columnconfigure(i,weight=(3,1,2)[i]);ttk.Label(m,text=lab).grid(row=0,column=i,sticky="w");ttk.Entry(m,textvariable=var).grid(row=1,column=i,sticky="ew",padx=(0 if i==0 else 8,0))
        a=ttk.LabelFrame(o,text="Agregar servidor",padding=10);a.pack(fill="x",pady=10)
        for i,(lab,var) in enumerate((("Servidor o equipo",self.server),("Fecha y hora del respaldo",self.time),("Observaciones",self.note))):a.columnconfigure(i,weight=2);ttk.Label(a,text=lab).grid(row=0,column=i,sticky="w");ttk.Entry(a,textvariable=var).grid(row=1,column=i,sticky="ew",padx=(0 if i==0 else 8,0))
        ttk.Button(a,text="Agregar",command=self.add).grid(row=1,column=3,padx=8)
        cols=("server","time","status","missing","notes");self.tree=ttk.Treeview(o,columns=cols,show="headings",selectmode="extended")
        for c,h,w in zip(cols,("Servidor","Último respaldo","Estado","Respaldos no realizados","Observaciones"),(220,170,120,350,190)):self.tree.heading(c,text=h);self.tree.column(c,width=w,anchor="center" if c in ("time","status") else "w")
        self.tree.pack(fill="both",expand=True);b=ttk.Frame(o);b.pack(fill="x",pady=10);ttk.Button(b,text="Indicar no realizados",command=self.missing).pack(side="left");ttk.Button(b,text="Eliminar",command=self.delete).pack(side="left",padx=6);ttk.Button(b,text="Guardar datos",command=self.save).pack(side="left",padx=(15,6));ttk.Button(b,text="Cargar datos",command=self.load).pack(side="left");ttk.Checkbutton(b,text="Word",variable=self.word).pack(side="right");ttk.Checkbutton(b,text="PDF",variable=self.pdf).pack(side="right");ttk.Button(b,text="Generar reporte",command=self.generate).pack(side="right",padx=12)
        import tkinter;self.general=tkinter.Text(o,height=3);self.general.pack(fill="x")
    def rows(self):return [ServerBackup(v[0],v[1],v[3],v[4]) for x in self.tree.get_children() for v in [self.tree.item(x,"values")]]
    def add(self):
        if not self.server.get().strip():messagebox.showwarning(APP_NAME,"Escribe el servidor.");return
        self.insert(ServerBackup(self.server.get().strip(),self.time.get().strip(),"",self.note.get().strip()));self.server.set("");self.note.set("")
    def insert(self,r):self.tree.insert("",END,values=(r.server,r.backup_time,r.status,r.missing,r.notes))
    def one(self):
        s=self.tree.selection()
        if len(s)!=1:messagebox.showinfo(APP_NAME,"Selecciona un servidor.");return None
        return s[0]
    def missing(self):
        x=self.one()
        if not x:return
        v=list(self.tree.item(x,"values"));z=simpledialog.askstring("No realizados","Escribe solo los nombres que no se hicieron, uno por línea. Déjalo vacío si todo salió bien.",initialvalue=v[3],parent=self.root)
        if z is not None:v[3]=z.strip();v[2]="Con incidencias" if z.strip() else "Realizado";self.tree.item(x,values=v)
    def delete(self):
        for x in self.tree.selection():self.tree.delete(x)
    def save(self):
        p=filedialog.asksaveasfilename(defaultextension=".json")
        if p:Path(p).write_text(json.dumps({"title":self.title.get(),"date":self.date.get(),"resp":self.resp.get(),"notes":self.general.get("1.0",END).strip(),"rows":[asdict(r) for r in self.rows()]},ensure_ascii=False,indent=2),encoding="utf-8")
    def load(self):
        p=filedialog.askopenfilename(filetypes=[("JSON","*.json")])
        if not p:return
        d=json.loads(Path(p).read_text(encoding="utf-8"));self.tree.delete(*self.tree.get_children());[self.insert(ServerBackup(**r)) for r in d["rows"]];self.title.set(d.get("title",""));self.date.set(d.get("date",""));self.resp.set(d.get("resp",""));self.general.delete("1.0",END);self.general.insert("1.0",d.get("notes",""))
    def generate(self):
        rows=self.rows();folder=filedialog.askdirectory()
        if not rows or not folder:return
        stem=safe_name(self.date.get().replace("/","_"))+"_reporte_respaldos";notes=self.general.get("1.0",END).strip()
        if self.word.get():generate_docx(Path(folder)/(stem+".docx"),self.title.get(),self.date.get(),self.resp.get(),rows,notes)
        if self.pdf.get():generate_pdf(Path(folder)/(stem+".pdf"),self.title.get(),self.date.get(),self.resp.get(),rows,notes)
        messagebox.showinfo(APP_NAME,"Reporte generado correctamente.");os.startfile(folder)

def main():r=Tk();App(r);r.mainloop()
if __name__=="__main__":main()
